"""Parser-backed document primitives (PDF, DOCX, PPTX, XLSX, CSV).

The LLM reasons over extracted text/tables. This is not a document agent.
"""

from __future__ import annotations

import csv
import os
import re
from pathlib import Path
from typing import Any

_MAX_DEFAULT = 80_000
_SUPPORTED = {".pdf", ".doc", ".docx", ".pptx", ".xlsx", ".csv", ".txt", ".md"}
SUPPORTED_EXTENSIONS = _SUPPORTED
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_UTF16_RUN = re.compile(rb"(?:[\x20-\x7e]\x00){12,}")
_DOC_NOISE = re.compile(
    r"^(Times New Roman|Calibri|Cambria|Arial|Microsoft|Normal\.dot|"
    r"Heading \d|Title|Caption|Style|Default Paragraph Font)$",
    re.I,
)


def _fail(code: str, message: str) -> dict[str, Any]:
    return {"status": "error", "error_code": code, "message": message}


def _clip(text: str, max_chars: int) -> tuple[str, bool]:
    if max_chars <= 0 or len(text) <= max_chars:
        return text, False
    return text[:max_chars] + f"\n...[truncated {len(text) - max_chars} chars]", True


def _extract_pdf(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return _fail("PARSER_UNAVAILABLE", "PyMuPDF is not installed. pip install pymupdf")
    try:
        doc = fitz.open(path)
    except Exception as exc:
        return _fail("PARSE_FAILED", f"Could not open PDF {path.name}: {exc}")
    try:
        pages = []
        total = 0
        truncated = False
        for i, page in enumerate(doc):
            if total >= max_chars:
                truncated = True
                break
            text = page.get_text("text") or ""
            remain = max_chars - total
            if len(text) > remain:
                text = text[:remain]
                truncated = True
            pages.append({"page": i + 1, "text": text})
            total += len(text)
        return {
            "status": "success",
            "format": "pdf",
            "path": str(path),
            "page_count": doc.page_count,
            "pages": pages,
            "text": "\n\n".join(p["text"] for p in pages if p["text"]),
            "truncated": truncated,
        }
    finally:
        doc.close()


def _extract_docx(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        from docx import Document
    except ImportError:
        return _fail("PARSER_UNAVAILABLE", "python-docx is not installed. pip install python-docx")
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    tables = []
    for table in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        tables.append(rows)
    text, truncated = _clip("\n".join(paragraphs), max_chars)
    return {
        "status": "success",
        "format": "docx",
        "path": str(path),
        "paragraphs": paragraphs[:500],
        "tables": tables[:50],
        "text": text,
        "truncated": truncated,
    }


def _extract_pptx(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        from pptx import Presentation
    except ImportError:
        return _fail("PARSER_UNAVAILABLE", "python-pptx is not installed. pip install python-pptx")
    pres = Presentation(str(path))
    slides = []
    chunks: list[str] = []
    for i, slide in enumerate(pres.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
        slides.append({"slide": i, "text": "\n".join(texts)})
        chunks.append(f"[slide {i}]\n" + "\n".join(texts))
    text, truncated = _clip("\n\n".join(chunks), max_chars)
    return {
        "status": "success",
        "format": "pptx",
        "path": str(path),
        "slide_count": len(slides),
        "slides": slides,
        "text": text,
        "truncated": truncated,
    }


def _extract_xlsx(path: Path, max_chars: int) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return _fail("PARSER_UNAVAILABLE", "openpyxl is not installed. pip install openpyxl")
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        sheets = []
        used = 0
        truncated = False
        for name in wb.sheetnames:
            ws = wb[name]
            rows: list[list[Any]] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    truncated = True
                    break
                values = ["" if c is None else c for c in row]
                rows.append(values)
                used += sum(len(str(v)) for v in values)
                if used >= max_chars:
                    truncated = True
                    break
            sheets.append({"name": name, "rows": rows, "row_count": len(rows)})
            if used >= max_chars:
                break
        plain_parts = []
        for sheet in sheets:
            plain_parts.append(f"[sheet {sheet['name']}]")
            for row in sheet["rows"]:
                plain_parts.append("\t".join(str(c) for c in row))
        text, clip_trunc = _clip("\n".join(plain_parts), max_chars)
        return {
            "status": "success",
            "format": "xlsx",
            "path": str(path),
            "sheet_names": list(wb.sheetnames),
            "sheets": sheets,
            "text": text,
            "truncated": truncated or clip_trunc,
        }
    finally:
        wb.close()


def _extract_csv(path: Path, max_chars: int) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as fh:
        sample = fh.read(4096)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample or "a,b\n1,2\n")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(fh, dialect)
        rows = []
        used = 0
        truncated = False
        for i, row in enumerate(reader):
            if i >= 500:
                truncated = True
                break
            rows.append(row)
            used += sum(len(str(c)) for c in row)
            if used >= max_chars:
                truncated = True
                break
    headers = rows[0] if rows else []
    plain = "\n".join(",".join(str(c) for c in row) for row in rows)
    text, clip_trunc = _clip(plain, max_chars)
    return {
        "status": "success",
        "format": "csv",
        "path": str(path),
        "headers": headers,
        "rows": rows,
        "row_count": len(rows),
        "text": text,
        "truncated": truncated or clip_trunc,
    }


def _rtf_to_text(raw: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    text = re.sub(r"\\[a-zA-Z]+\d* ?", " ", text)
    text = re.sub(r"[{}]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _ole_unicode_strings(data: bytes) -> str:
    chunks: list[str] = []
    seen: set[str] = set()
    for match in _UTF16_RUN.finditer(data):
        piece = match.group().decode("utf-16-le", errors="ignore").strip()
        if len(piece) < 12 or _DOC_NOISE.match(piece) or piece in seen:
            continue
        seen.add(piece)
        chunks.append(piece)
    return "\n".join(chunks)


def _extract_doc_winword(path: Path, max_chars: int) -> dict[str, Any]:
    """Last-resort Word COM on Windows. Optional — heuristic usually suffices."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return _fail(
            "PARSER_UNAVAILABLE",
            f"{path.name} is a legacy .doc. Save it as .docx and attach that.",
        )
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(path), ReadOnly=True, AddToRecentFiles=False)
        try:
            raw = str(doc.Content.Text or "").replace("\r", "\n")
        finally:
            doc.Close(False)
        text, truncated = _clip(raw, max_chars)
        if not text.strip():
            return _fail(
                "PARSE_FAILED",
                f"{path.name} has no extractable text. Save as .docx if you can.",
            )
        return {
            "status": "success",
            "format": "doc",
            "path": str(path),
            "text": text,
            "truncated": truncated,
        }
    except Exception as exc:
        return _fail(
            "PARSE_FAILED",
            f"Could not read {path.name} ({exc}). Save it as .docx and attach that.",
        )
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _extract_legacy_doc(path: Path, max_chars: int) -> dict[str, Any]:
    """Word 97-2003 .doc (OLE). python-docx cannot read this format."""
    data = path.read_bytes()
    if data[:4] == b"%PDF":
        return _extract_pdf(path, max_chars)
    if data[:2] == b"PK":
        return _extract_docx(path, max_chars)
    stripped = data.lstrip()
    if stripped.startswith(b"{\\rtf"):
        text, truncated = _clip(
            _rtf_to_text(stripped.decode("latin-1", errors="replace")), max_chars
        )
        return {
            "status": "success",
            "format": "doc",
            "path": str(path),
            "text": text,
            "truncated": truncated,
        }
    text = _ole_unicode_strings(data)
    if len(text) < 80 and data[:8] == _OLE_MAGIC:
        com = _extract_doc_winword(path, max_chars)
        if com.get("status") == "success":
            return com
    if not text.strip():
        return _fail(
            "PARSE_FAILED",
            f"{path.name} is a legacy .doc with no extractable text. "
            "Save it as .docx in Word and attach that.",
        )
    text, truncated = _clip(text, max_chars)
    return {
        "status": "success",
        "format": "doc",
        "path": str(path),
        "text": text,
        "truncated": truncated,
    }


def extract_document(path: str, max_chars: int = _MAX_DEFAULT) -> dict[str, Any]:
    """Extract structured text/tables from a document file."""
    if not path or not str(path).strip():
        return _fail("INVALID_ARGS", "extract_document requires a path.")
    target = Path(os.path.expanduser(str(path))).resolve()
    if not target.is_file():
        return _fail("NOT_FOUND", f"File not found: {target}")
    try:
        max_c = int(max_chars)
    except (TypeError, ValueError):
        max_c = _MAX_DEFAULT
    max_c = max(1_000, min(max_c, 400_000))

    suffix = target.suffix.lower()
    try:
        with target.open("rb") as fh:
            head = fh.read(16)
    except OSError as exc:
        return _fail("NOT_FOUND", f"File not found: {target} ({exc})")
    if head.startswith(b"%PDF"):
        suffix = ".pdf"
    elif suffix == ".doc" and head.startswith(b"PK"):
        suffix = ".docx"
    elif suffix == ".docx" and head[:8] == _OLE_MAGIC:
        suffix = ".doc"

    try:
        if suffix == ".pdf":
            return _extract_pdf(target, max_c)
        if suffix == ".doc":
            return _extract_legacy_doc(target, max_c)
        if suffix == ".docx":
            return _extract_docx(target, max_c)
        if suffix == ".pptx":
            return _extract_pptx(target, max_c)
        if suffix == ".xlsx":
            return _extract_xlsx(target, max_c)
        if suffix == ".csv":
            return _extract_csv(target, max_c)
        if suffix in {".txt", ".md"}:
            text, truncated = _clip(
                target.read_text(encoding="utf-8", errors="replace"), max_c
            )
            return {
                "status": "success",
                "format": suffix.lstrip("."),
                "path": str(target),
                "text": text,
                "truncated": truncated,
            }
    except Exception as exc:
        return _fail("PARSE_FAILED", f"Could not parse {target.name}: {exc}")
    return _fail(
        "UNSUPPORTED_FORMAT",
        f"Unsupported format {suffix or '(none)'}. Supported: {sorted(_SUPPORTED)}",
    )


def extracted_text(result: dict[str, Any]) -> str:
    """Plain text for prompts. CSV/XLSX now include a ``text`` field; this is the fallback."""
    if not isinstance(result, dict) or result.get("status") != "success":
        return ""
    text = result.get("text")
    if isinstance(text, str) and text.strip():
        return text
    if result.get("format") == "xlsx":
        lines: list[str] = []
        for sheet in result.get("sheets") or []:
            lines.append(f"[sheet {sheet.get('name')}]")
            for row in sheet.get("rows") or []:
                lines.append("\t".join(str(c) for c in row))
        return "\n".join(lines)
    if result.get("format") == "csv":
        return "\n".join(
            ",".join(str(c) for c in row) for row in (result.get("rows") or [])
        )
    pages = result.get("pages") or []
    if pages:
        return "\n\n".join(str(p.get("text") or "") for p in pages)
    return ""
